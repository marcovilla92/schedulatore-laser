"""XLSX preventivo (quote) exporter.

Extracted from preventivatore 2.0.py — pure function, no UI references.
"""

import logging
import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

logger = logging.getLogger(__name__)


def esporta_xlsx_preventivo(
    path: str,
    articoli_processati: list,
    config: dict,
    cliente: str,
    ordine: str,
    margine: float,
    costi_montaggio: dict,
    pdf_order_codes: list,
) -> str | None:
    """Esporta il preventivo corrente in un file XLSX semplificato per il cliente.

    Args:
        path: percorso di output del file XLSX.
        articoli_processati: lista di articoli processati con costi.
        config: dizionario di configurazione.
        cliente: nome del cliente.
        ordine: numero ordine.
        margine: percentuale di margine (es. 15.0 per 15%).
        costi_montaggio: dict {codice_assieme: dati_montaggio} con costi, ore, ecc.
        pdf_order_codes: lista di codici dal PDF ordine per ordinare l'export.

    Returns:
        path su successo, None su fallimento.
    """
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Preventivo"

        # Header: Cliente, Ordine, Data
        data_oggi = datetime.now().strftime('%d/%m/%Y')

        ws["A1"] = "PREVENTIVO"
        ws["A1"].font = Font(size=18, bold=True)
        ws["A2"] = f"Cliente: {cliente}"
        ws["A2"].font = Font(size=12, bold=True)
        ws["A3"] = f"Numero Ordine: {ordine}"
        ws["A3"].font = Font(size=12)
        ws["A4"] = f"Data: {data_oggi}"
        ws["A4"].font = Font(size=12)
        ws["A5"] = ""  # Riga vuota

        # Tabella articoli - Header in riga 6
        row_start = 6
        ws[f"A{row_start}"] = "Codice"
        ws[f"B{row_start}"] = "Prezzo Unitario"
        ws[f"C{row_start}"] = "Quantità"
        ws[f"D{row_start}"] = "Totale"

        # Formatta header
        for col in ['A', 'B', 'C', 'D']:
            cell = ws[f"{col}{row_start}"]
            cell.font = Font(bold=True, size=11)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, size=11, color="FFFFFF")
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Ordina articoli secondo l'ordine del PDF (se caricato)
        articoli_da_esportare = articoli_processati[:]

        if pdf_order_codes:
            # Crea dizionario codice -> posizione nel PDF
            pdf_positions = {code: idx for idx, code in enumerate(pdf_order_codes)}

            # Separa articoli: quelli nel PDF e quelli non nel PDF
            articoli_nel_pdf = []
            articoli_non_nel_pdf = []

            for art in articoli_da_esportare:
                codice = art.get('codice', '')
                codice_assieme = art.get('codice_assieme', None)

                # Cerca nel PDF sia il codice articolo che il codice assieme
                if codice_assieme and codice_assieme in pdf_positions:
                    # Usa la posizione del codice assieme
                    articoli_nel_pdf.append((pdf_positions[codice_assieme], art))
                elif codice in pdf_positions:
                    # Usa la posizione del codice articolo
                    articoli_nel_pdf.append((pdf_positions[codice], art))
                else:
                    articoli_non_nel_pdf.append(art)

            # Ordina gli articoli nel PDF per posizione
            articoli_nel_pdf.sort(key=lambda x: x[0])

            # Combina: prima quelli del PDF (ordinati), poi gli altri
            articoli_da_esportare = [art for pos, art in articoli_nel_pdf] + articoli_non_nel_pdf

        # Raggruppa articoli: assiemi in una riga (somma componenti + montaggio), sciolti singoli
        montaggio_data = costi_montaggio or {}

        # Calcola prezzo per ogni articolo
        def _prezzo_articolo(art):
            costo_mat = float(art.get('costo', 0.0) or 0.0)
            area_a = float(art.get('area', 0.0) or 0.0)
            c_piega = float(art.get('costo_piega', 0.0) or 0.0)
            c_sald = float(art.get('costo_saldatura', 0.0) or 0.0)
            c_filett = float(art.get('costo_filettatura', 0.0) or 0.0)
            c_svas = float(art.get('costo_svasatura', 0.0) or 0.0)
            costo = costo_mat + c_piega + c_sald + c_filett + c_svas
            if margine > 0:
                costo = costo * (1 + margine / 100)
            return costo

        # Separa assiemi da articoli sciolti
        assiemi_prezzi = {}  # codice_assieme -> prezzo totale (somma componenti * quantità)
        articoli_sciolti = []  # (codice, prezzo_unitario, quantita)

        for art in articoli_da_esportare:
            codice_assieme = art.get('codice_assieme', '') or ''
            prezzo_unitario = _prezzo_articolo(art)
            qty = art.get('quantita', 1)
            if codice_assieme:
                # Usa la qty del componente dal popup montaggio (se specificata)
                comp_qty_map = montaggio_data.get(codice_assieme, {}).get('componenti_qty', {})
                codice_art = art.get('codice', '')
                qty_comp = comp_qty_map.get(codice_art, qty)
                assiemi_prezzi[codice_assieme] = assiemi_prezzi.get(codice_assieme, 0.0) + (prezzo_unitario * qty_comp)
            else:
                articoli_sciolti.append((art.get('codice', ''), prezzo_unitario, qty))

        # Aggiungi costo montaggio agli assiemi
        for codice_ass, mont in montaggio_data.items():
            costo_mont = mont['costo']
            if margine > 0:
                costo_mont = costo_mont * (1 + margine / 100)
            assiemi_prezzi[codice_ass] = assiemi_prezzi.get(codice_ass, 0.0) + costo_mont

        # Scrivi righe
        current_row = row_start + 1

        # Prima gli assiemi (una riga per assieme, prezzo già sommato)
        for codice_ass, prezzo_totale in assiemi_prezzi.items():
            qty_ass = montaggio_data.get(codice_ass, {}).get('qty', 1)
            ws[f"A{current_row}"] = codice_ass
            ws[f"B{current_row}"] = prezzo_totale
            ws[f"C{current_row}"] = qty_ass
            ws[f"D{current_row}"] = f"=B{current_row}*C{current_row}"
            current_row += 1

        # Poi gli articoli sciolti (con quantità reale)
        for codice_art, prezzo, qty in articoli_sciolti:
            ws[f"A{current_row}"] = codice_art
            ws[f"B{current_row}"] = prezzo
            ws[f"C{current_row}"] = qty
            ws[f"D{current_row}"] = f"=B{current_row}*C{current_row}"
            current_row += 1

        # Riga totale
        total_row = current_row
        ws[f"A{total_row}"] = "TOTALE"
        ws[f"A{total_row}"].font = Font(bold=True, size=12)
        ws[f"D{total_row}"] = f"=SUM(D{row_start+1}:D{total_row-1})"
        ws[f"D{total_row}"].font = Font(bold=True, size=12)

        # Formatta colonne
        currency_fmt = '#,##0.00 €'
        int_fmt = '0'

        for row in range(row_start + 1, total_row + 1):
            ws[f"B{row}"].number_format = currency_fmt
            ws[f"D{row}"].number_format = currency_fmt
            ws[f"C{row}"].number_format = int_fmt

        # Larghezza colonne
        ws.column_dimensions['A'].width = 24  # Codice
        ws.column_dimensions['B'].width = 18  # Prezzo Unitario
        ws.column_dimensions['C'].width = 12  # Quantità
        ws.column_dimensions['D'].width = 18  # Totale

        # Aggiungi bordi alla tabella
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for row in range(row_start, total_row + 1):
            for col in ['A', 'B', 'C', 'D']:
                ws[f"{col}{row}"].border = thin_border

        # Aggiungi nota informativa
        note_row = total_row + 2
        ws[f"A{note_row}"] = "Nota: Modificare la quantità (colonna C) per aggiornare automaticamente il totale."
        ws[f"A{note_row}"].font = Font(italic=True, size=9, color="666666")
        ws.merge_cells(f"A{note_row}:D{note_row}")

        wb.save(path)

        return path

    except Exception as e:
        logger.error("Impossibile esportare in Excel: %s", e)
        return None
