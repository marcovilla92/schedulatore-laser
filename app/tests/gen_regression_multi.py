"""Ingestione MULTI-ORDINE per la regression suite del parsing.

Scansiona _test_input/lantek_orders/<ordine>/ dove ogni cartella-ordine
contiene:
  - UN file .xlsx = export Lantek (foglio 'Struttura' con Area/Perimetro/
    Materiale/Spessore/Peso per Codice)
  - N file .dxf/.DXF = i disegni dei pezzi (nome file = Codice)

Per ogni pezzo con DXF + riga Lantek, genera un caso in
_test_input/regression/<ordine>__<codice>/ con pezzo.dxf + expected.json
(valori Lantek come ground-truth). Poi:
    python app/tests/baseline_parsing.py --write   # trova i click e attiva
    python app/tests/test_regression_dxf.py        # misura

Uso: python app/tests/gen_regression_multi.py [--force]

Marco: metti gli ordini in app/../_test_input/lantek_orders/<nome_ordine>/
(un xlsx + i DXF). Ogni ordine in una sua cartella. I nomi DXF devono
combaciare col Codice nell'xlsx.
"""
from __future__ import annotations
import argparse
import json
import shutil
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERRORE: openpyxl non installato. pip install openpyxl", file=sys.stderr)
    sys.exit(2)

_APP = Path(__file__).resolve().parent.parent
_ROOT = _APP.parent
_ORDERS = _ROOT / "_test_input" / "lantek_orders"
_REG = _ROOT / "_test_input" / "regression"


def _norm_mat(raw: str) -> str:
    r = (raw or "").upper().strip()
    if "316" in r:
        return "INOX_316"
    if "INOX" in r or "AISI 304" in r or "1.4301" in r:
        return "INOX_304"
    if "5083" in r:
        return "ALU_5083"
    if "ALU" in r or "5754" in r:
        return "ALU_5754"
    if "ZINC" in r or "DX51" in r:
        return "ZINCATO"
    if "OTTON" in r:
        return "OTTONE"
    return "S235"


def _sanitize(s: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in s)


def _read_lantek_xlsx(xlsx_path: Path) -> dict:
    """Ritorna {codice: {area_dm2, perim_m, materiale, spessore_mm, peso_kg}}."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Struttura"] if "Struttura" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]

    def col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    ci = {
        "codice": col("Codice"),
        "area": col("Area"),
        "perim": col("Perimetro di taglio", "Perimetro"),
        "peso": col("Peso"),
        "mat": col("Materiale"),
        "sp": col("Spessore"),
    }
    if ci["codice"] is None or ci["area"] is None:
        return {}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        cod = row[ci["codice"]]
        if not cod:
            continue
        area_m2 = row[ci["area"]] if ci["area"] is not None else None
        out[str(cod).strip().upper()] = {
            "area_dm2": round(float(area_m2) * 100, 4) if area_m2 else None,  # m² → dm²
            "perim_m": round(float(row[ci["perim"]]), 4) if ci["perim"] is not None and row[ci["perim"]] else None,
            "peso_kg": round(float(row[ci["peso"]]), 4) if ci["peso"] is not None and row[ci["peso"]] else None,
            "materiale": _norm_mat(str(row[ci["mat"]] or "")) if ci["mat"] is not None else "S235",
            "spessore_mm": float(row[ci["sp"]]) if ci["sp"] is not None and row[ci["sp"]] else None,
        }
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    if not _ORDERS.is_dir():
        _ORDERS.mkdir(parents=True, exist_ok=True)
        print(f"Creata cartella {_ORDERS}")
        print("Metti dentro un sotto-cartella per ogni ordine (xlsx Lantek + DXF), poi rilancia.")
        return

    orders = [d for d in _ORDERS.iterdir() if d.is_dir()]
    if not orders:
        print(f"Nessun ordine in {_ORDERS}. Metti le cartelle-ordine e rilancia.")
        return

    tot_cases = tot_nodxf = tot_noarea = 0
    for order in sorted(orders):
        xlsxs = list(order.glob("*.xlsx"))
        if not xlsxs:
            print(f"[SKIP] {order.name}: nessun .xlsx")
            continue
        lantek = _read_lantek_xlsx(xlsxs[0])
        if not lantek:
            print(f"[SKIP] {order.name}: xlsx senza foglio Struttura/colonne")
            continue
        dxfs = {p.stem.upper(): p for p in order.rglob("*") if p.suffix.lower() in (".dxf",)}
        made = 0
        for cod, dxf in dxfs.items():
            info = lantek.get(cod)
            if not info:
                tot_nodxf += 1
                continue
            if not info.get("area_dm2"):
                tot_noarea += 1
                continue
            case_name = f"{_sanitize(order.name)}__{_sanitize(cod)}"
            case_dir = _REG / case_name
            if case_dir.exists() and not args.force:
                continue
            case_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(dxf, case_dir / "pezzo.dxf")
            (case_dir / "expected.json").write_text(json.dumps({
                "click_x_mm": None, "click_y_mm": None,
                "expected": {
                    "area_dm2": info["area_dm2"], "perim_m": info["perim_m"],
                    "n_fori": None, "materiale": info["materiale"],
                    "spessore_mm": info["spessore_mm"], "peso_kg": info["peso_kg"],
                },
                "tolleranze": {"area_dm2_pct": 0.5, "perim_m_pct": 1.5, "n_fori_abs": 0},
                "misurato_con": f"Lantek export {xlsxs[0].name} (ordine {order.name})",
            }, indent=2, ensure_ascii=False), encoding="utf-8")
            made += 1
            tot_cases += 1
        print(f"[OK] {order.name}: {made} casi generati ({len(dxfs)} DXF, {len(lantek)} righe Lantek)")

    print(f"\nTotale casi generati: {tot_cases}  (DXF senza riga Lantek: {tot_nodxf}, righe senza area: {tot_noarea})")
    if tot_cases:
        print("Ora: python app/tests/baseline_parsing.py --write  (trova i click e attiva)")
        print("Poi: python app/tests/test_regression_dxf.py       (misura l'accuratezza)")


if __name__ == "__main__":
    main()
